"""Exportación a Excel, simulación de cierres y las consultas como herramientas de LangChain."""

from __future__ import annotations

from io import BytesIO

import pytest
from apiutil import DS, EMP, MARZO
from openpyxl import load_workbook

from tempo_ledger.api.service import NotFound
from tempo_ledger.api.tools import build_tools
from tempo_ledger.domain import Period
from tempo_ledger.pipeline import ClosingRequest, build_closing_chain


@pytest.fixture(scope="module")
def report():
    return build_closing_chain(lambda: DS).invoke(ClosingRequest(Period.parse(MARZO)))


# --- Excel -------------------------------------------------------------------------------

def test_el_excel_trae_los_mismos_numeros_que_la_api(client, report):
    r = client.get(f"/closings/{MARZO}/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert f'filename="cierre-{MARZO}.xlsx"' in r.headers["content-disposition"]

    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Resumen", f"Cierre {MARZO}", "Líneas"]
    filas = list(wb[f"Cierre {MARZO}"].iter_rows(min_row=2, values_only=True))
    api = client.get(f"/closings/{MARZO}").json()["rows"]
    assert len(filas) == len(api)
    for fila, row in zip(filas, api, strict=True):
        assert fila[0] == row["employee_id"]
        assert fila[3] == row["closing"]["status_label"]
        assert fila[10] == round(row["closing"]["closing"] / 60, 2)  # saldo final (h)
    lineas = list(wb["Líneas"].iter_rows(min_row=2, values_only=True))
    assert len(lineas) == sum(len(r.closing.lines) for r in report.rows)


def test_el_excel_de_un_cierre_inexistente(client):
    assert client.get("/closings/2030-01/export.xlsx").status_code == 404


# --- Simulación --------------------------------------------------------------------------

def test_simular_con_el_calendario_real_no_cambia_nada(client):
    s = client.post("/simulations", json={"period": MARZO}).json()
    real = client.get(f"/closings/{MARZO}").json()
    assert s["changes"] == []
    assert s["simulated"]["total_deduction"] == real["total_deduction"]
    assert s["baseline_by_status"] == real["by_status"]


def test_simular_otro_corte_no_guarda_nada(client):
    antes = client.get(f"/closings/{MARZO}").json()
    s = client.post("/simulations", json={"period": MARZO, "cutoff_day": 15,
                                          "effective_business_days": 1}).json()
    assert (s["simulated"]["cutoff"], s["simulated"]["effective"]) == ("2026-03-13",
                                                                        "2026-03-16")
    assert s["changes"]  # cambiar el corte mueve asientos entre cierres
    assert client.get(f"/closings/{MARZO}").json() == antes


def test_simular_otro_vencimiento(client):
    s = client.post("/simulations", json={"period": MARZO, "debit_expiry_days": 0}).json()
    assert s["debit_expiry_days"] == 0
    assert s["simulated"]["total_deduction"] != s["baseline_total_deduction"]


@pytest.mark.parametrize("cuerpo, codigo", [
    ({"period": MARZO, "cutoff_day": 0}, 422),
    ({"period": MARZO, "effective_business_days": -1}, 422),
    ({"period": "2030-01"}, 404),
    ({"period": "marzo"}, 422),
])
def test_simulaciones_invalidas(client, cuerpo, codigo):
    assert client.post("/simulations", json=cuerpo).status_code == codigo


# --- Herramientas de LangChain -----------------------------------------------------------

@pytest.fixture(scope="module")
def tools(client):
    return {t.name: t for t in build_tools(client.app.state.service)}


def test_las_herramientas_tienen_nombre_descripcion_y_argumentos(tools):
    assert set(tools) == {"saldo_del_cierre", "eventos_del_cierre", "serie_del_saldo"}
    for t in tools.values():
        assert t.description
    assert set(tools["saldo_del_cierre"].args) == {"employee_id", "period"}


def test_cada_herramienta_devuelve_lo_mismo_que_la_api(client, tools):
    assert tools["saldo_del_cierre"].invoke({"employee_id": EMP, "period": MARZO}) == (
        client.get(f"/employees/{EMP}/closings/{MARZO}").json())
    assert tools["serie_del_saldo"].invoke({"employee_id": EMP, "until": MARZO}) == (
        client.get(f"/employees/{EMP}/series", params={"until": MARZO}).json())
    tardias = tools["eventos_del_cierre"].invoke({"period": MARZO, "light": "fuera_de_plazo"})
    api = client.get(f"/closings/{MARZO}/events").json()["events"]
    assert tardias == [e for e in api if e["light"] == "fuera_de_plazo"]


def test_las_herramientas_validan_sus_argumentos(tools):
    with pytest.raises(ValueError):
        tools["saldo_del_cierre"].invoke({"employee_id": "uno", "period": MARZO})
    with pytest.raises(NotFound):
        tools["saldo_del_cierre"].invoke({"employee_id": 999, "period": MARZO})
