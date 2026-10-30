"""Los comandos `tempo-ledger seed`, `export` y `openapi`."""

from __future__ import annotations

import json
from pathlib import Path

from dbutil import database_url
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from tempo_ledger.cli import main


def test_seed_llena_la_base(tmp_path: Path, capsys) -> None:
    url = database_url(tmp_path)
    main(["seed", "--seed", "7", "--employees", "15", "--url", url])
    engine = create_engine(url)
    try:
        with engine.connect() as c:
            assert c.execute(text("select count(*) from employees")).scalar() == 15
    finally:
        engine.dispose()
    assert "base cargada: 15 empleados" in capsys.readouterr().out


def test_export_escribe_la_muestra(tmp_path: Path) -> None:
    out = tmp_path / "muestra.xlsx"
    main(["export", "--seed", "7", "--employees", "15", "--out", str(out)])
    assert "Resumen" in load_workbook(out).sheetnames


def test_openapi_escribe_el_contrato(tmp_path: Path) -> None:
    out = tmp_path / "openapi.json"
    main(["openapi", "--out", str(out)])
    contrato = json.loads(out.read_text(encoding="utf-8"))
    assert contrato["info"]["title"] == "tempo-ledger"
    assert "/closings/{period}" in contrato["paths"]


def test_el_contrato_de_la_web_esta_al_dia(tmp_path: Path) -> None:
    """Si cambia la API, hay que correr `tempo-ledger openapi` y `npm run types`."""
    out = tmp_path / "openapi.json"
    main(["openapi", "--out", str(out)])
    guardado = Path(__file__).parents[1] / "frontend" / "openapi.json"
    assert json.loads(out.read_text(encoding="utf-8")) == json.loads(
        guardado.read_text(encoding="utf-8"))
