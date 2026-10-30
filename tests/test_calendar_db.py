"""Calendario de cierres y persistencia (migración + ida y vuelta en SQLite)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from dbutil import database_url
from sqlalchemy import inspect

from tempo_ledger.data.export import export_xlsx
from tempo_ledger.data.generator import generate
from tempo_ledger.domain.calendar import (
    add_business_days,
    closing_period,
    cutoff_date,
    effective_date,
)
from tempo_ledger.domain.models import Period
from tempo_ledger.infra import db, load


def test_corte_retrocede_al_dia_habil_anterior() -> None:
    assert cutoff_date(Period(2026, 10)) == date(2026, 10, 23)  # 25 cae domingo
    assert cutoff_date(Period(2026, 12)) == date(2026, 12, 24)  # 25 es feriado
    assert cutoff_date(Period(2026, 9)) == date(2026, 9, 25)


def test_fecha_efectiva_tres_dias_habiles_despues() -> None:
    assert effective_date(Period(2026, 9)) == date(2026, 9, 30)
    assert effective_date(Period(2026, 10)) == date(2026, 10, 28)
    assert add_business_days(date(2026, 7, 27), 1) == date(2026, 7, 30)  # 28 y 29 feriados


def test_periodo_de_cierre_sigue_la_ventana() -> None:
    assert closing_period(date(2026, 9, 25)) == Period(2026, 9)
    assert closing_period(date(2026, 9, 26)) == Period(2026, 10)
    assert closing_period(date(2026, 12, 28)) == Period(2027, 1)




def test_migracion_crea_las_tablas(tmp_path: Path) -> None:
    url = database_url(tmp_path)
    db.upgrade(url)
    engine = db.make_engine(url)
    try:
        tablas = set(inspect(engine).get_table_names())
    finally:
        engine.dispose()
    assert {"employees", "requests", "attendance", "closings", "exceptions"} <= tablas


def test_ida_y_vuelta_por_la_base(tmp_path: Path) -> None:
    url = database_url(tmp_path)
    db.upgrade(url)
    engine = db.make_engine(url)
    ds = generate(seed=3, employees=15)
    try:
        load.save(engine, ds)
        load.save(engine, ds)  # recargar reemplaza, no duplica
        back = load.read(engine)
    finally:
        engine.dispose()
    assert back.employees == ds.employees
    assert back.closings == ds.closings
    assert back.exceptions == ds.exceptions
    assert sorted(back.requests, key=lambda r: r.id) == sorted(ds.requests, key=lambda r: r.id)
    assert back.attendance == sorted(ds.attendance, key=lambda a: (a.employee_id, a.day))


def test_exporta_la_muestra(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    path = export_xlsx(generate(seed=3, employees=15), tmp_path / "muestra.xlsx")
    wb = load_workbook(path, read_only=True)
    assert wb.sheetnames == ["Resumen", "Empleados", "Cierres", "Solicitudes", "Asistencia",
                             "Excepciones"]
