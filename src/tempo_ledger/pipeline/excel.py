"""El cierre de todos los empleados en Excel: resumen, saldos y las líneas que los explican.

Los tiempos van en **horas decimales** (1,50 = 1h30) para que se puedan sumar y ordenar;
el cálculo sigue siendo en minutos enteros.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from tempo_ledger.data.export import autowidth, write_table
from tempo_ledger.domain.ledger import Effect
from tempo_ledger.pipeline.closing import ClosingReport

EFFECTS = {Effect.CREDIT: "suma", Effect.DUE: "vence", Effect.PENDING: "por vencer",
           Effect.SETTLED: "se liquida en el finiquito"}
HOURS_FORMAT = "0.00"


def hours(minutes: int | None) -> float | None:
    return None if minutes is None else round(minutes / 60, 2)


def closing_workbook(report: ClosingReport) -> bytes:
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    summary.title = "Resumen"
    rows: list[list[object]] = [
        [f"Cierre {report.period} — tempo-ledger (datos ficticios)"],
        [],
        ["fecha de corte", report.cutoff],
        ["fecha efectiva", report.effective],
        ["empleados", len(report.rows)],
        ["descuento total (h)", hours(report.total_deduction)],
        ["omitido por excepción (h)", hours(report.total_waived)],
        [],
        ["estado", "empleados"],
        *([status, n] for status, n in sorted(report.by_status.items())),
        [],
        ["Horas decimales: 1,50 = 1h30. Las columnas (h) se pueden sumar y ordenar."],
    ]
    for row in rows:
        summary.append(row)
    summary["A1"].font = Font(bold=True, size=13)
    for cell in (summary["B3"], summary["B4"]):
        cell.number_format = "yyyy-mm-dd"
    autowidth(summary)

    balances = wb.create_sheet(f"Cierre {report.period}")
    write_table(
        balances,
        ["id", "empleado", "equipo", "estado", "saldo anterior (h)", "a favor (h)",
         "vencen (h)", "saldo (h)", "descuento (h)", "omitido por excepción (h)",
         "saldo final (h)", "por vencer (h)", "tardanza no recuperable (h)", "finiquito (h)",
         "explicación"],
        ([r.employee_id, r.full_name, r.team, r.closing.status.label,
          *(hours(m) for m in (r.closing.opening, r.closing.credits, r.closing.due,
                               r.closing.balance, r.closing.deduction, r.closing.waived,
                               r.closing.closing, r.closing.pending,
                               r.closing.unrecoverable_late, r.closing.settlement)),
          r.closing.explain()] for r in report.rows),
    )
    for column in balances.iter_cols(min_col=5, max_col=14, min_row=2):
        for cell in column:
            cell.number_format = HOURS_FORMAT
    balances.column_dimensions["O"].width = 90

    lines = wb.create_sheet("Líneas")
    write_table(
        lines,
        ["id", "empleado", "solicitud", "tipo", "fecha", "horas", "efecto", "afecta el",
         "fuera de plazo"],
        ([r.employee_id, r.full_name, ln.request_id, ln.kind.value, ln.day, hours(ln.minutes),
          EFFECTS[ln.effect], ln.due_on, "sí" if ln.late else "no"]
         for r in report.rows for ln in r.closing.lines),
    )
    for cell in lines["F"][1:]:
        cell.number_format = HOURS_FORMAT

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
