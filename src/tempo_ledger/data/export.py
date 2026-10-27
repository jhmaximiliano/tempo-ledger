"""Muestra del conjunto ficticio en Excel, para revisarla a ojo (punto de control de F1).

Solo presenta los datos: no aplica las reglas del libro. Las columnas «tarde (min)» y
«almuerzo (min)» de Asistencia son ayudas de lectura, calculadas de las marcas.
"""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from tempo_ledger.data.generator import EDGE_CASES, Dataset

_HEADER_FILL = PatternFill("solid", fgColor="DDE4EE")
_THIN = Side(style="thin", color="B7C0CC")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_EDGE_NOTES = {
    "siempre_a_favor": "perfil acumulador: muchas horas extra, pocos permisos",
    "deuda_grande": "perfil deudor con el doble de permisos: saldos en contra",
    "aprobacion_tardia": "sus aprobaciones llegan después de la fecha efectiva del cierre",
    "cese_mitad_de_mes": "cesa el 2026-06-15: sin asistencia ni solicitudes después",
    "ingreso_reciente": "ingresa el 2026-09-14: historia corta",
    "sin_movimientos": "ninguna solicitud en todo el periodo (solo asistencia)",
    "excepcion_parcial": "permiso de 2 días en marzo-2026 con excepción parcial en 2026-04",
}


def export_xlsx(ds: Dataset, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    names = {e.id: e.full_name for e in ds.employees}
    edge_of = {v: k for k, v in ds.edge_cases.items()}
    wb = Workbook()

    ws = wb.active
    assert ws is not None
    ws.title = "Resumen"
    _summary(ws, ds, names)

    _table(wb.create_sheet("Empleados"),
           ["id", "nombre", "equipo", "modalidad", "sede", "ingreso", "cese", "perfil",
            "caso borde"],
           ([e.id, e.full_name, e.team, e.modality.value, e.site, e.hired_on, e.terminated_on,
             ds.profiles.get(e.id, ""), edge_of.get(e.id, "")] for e in ds.employees))
    _table(wb.create_sheet("Cierres"), ["periodo", "fecha de corte", "fecha efectiva"],
           ([c.period, c.cutoff, c.effective] for c in ds.closings))
    _table(wb.create_sheet("Solicitudes"),
           ["id", "empleado", "nombre", "tipo", "estado", "desde", "hasta", "minutos", "horas",
            "registrada", "resuelta", "motivo"],
           ([r.id, r.employee_id, names[r.employee_id], r.kind.value, r.status.value,
             r.start_date, r.end_date, r.minutes, round(r.minutes / 60, 2), r.registered_at,
             r.decided_at, r.reason] for r in ds.requests))
    _table(wb.create_sheet("Asistencia"),
           ["empleado", "nombre", "día", "entrada", "sale almuerzo", "vuelve almuerzo",
            "salida", "tarde (min)", "almuerzo (min)"],
           ([a.employee_id, names[a.employee_id], a.day, a.check_in, a.lunch_out, a.lunch_in,
             a.check_out, _late(a.check_in), _span(a.lunch_out, a.lunch_in)]
            for a in ds.attendance))
    _table(wb.create_sheet("Excepciones"),
           ["periodo", "empleado", "nombre", "tipo", "minutos", "motivo"],
           ([x.period, x.employee_id, names[x.employee_id], x.kind.value, x.minutes, x.reason]
            for x in ds.exceptions))
    wb.save(path)
    return path


def _minutes(t: time) -> int:
    return t.hour * 60 + t.minute


def _late(check_in: time | None) -> int | None:
    return max(0, _minutes(check_in) - 9 * 60) if check_in else None


def _span(a: time | None, b: time | None) -> int | None:
    return _minutes(b) - _minutes(a) if a and b else None


def _summary(ws: Worksheet, ds: Dataset, names: dict[int, str]) -> None:
    rows: list[list[object]] = [
        ["Muestra de datos ficticios — tempo-ledger"],
        [],
        ["semilla", ds.seed],
        ["empleados", len(ds.employees)],
        ["ingresos durante el periodo", sum(e.hired_on.year >= 2025 for e in ds.employees)],
        ["ceses durante el periodo", sum(e.terminated_on is not None for e in ds.employees)],
        ["cierres", f"{len(ds.closings)} ({ds.closings[0].period} a {ds.closings[-1].period})"],
        ["solicitudes", len(ds.requests)],
        ["marcas de asistencia (días)", len(ds.attendance)],
        ["excepciones", len(ds.exceptions)],
        [],
        ["Solicitudes por tipo y estado"],
    ]
    by_kind = Counter((r.kind.value, r.status.value) for r in ds.requests)
    statuses = sorted({s for _, s in by_kind})
    rows.append(["tipo", *statuses, "total"])
    for kind in sorted({k for k, _ in by_kind}):
        counts = [by_kind[(kind, s)] for s in statuses]
        rows.append([kind, *counts, sum(counts)])
    rows += [[], ["Perfiles de comportamiento"]]
    rows += [[p, n] for p, n in sorted(Counter(ds.profiles.values()).items())]
    rows += [[], ["Casos borde (qué mirar)"], ["caso", "empleado", "nombre", "qué esperar"]]
    rows += [[c, ds.edge_cases[c], names[ds.edge_cases[c]], _EDGE_NOTES[c]] for c in EDGE_CASES]
    for row in rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=13)
    for cell in ws["A"]:
        if cell.value in ("Solicitudes por tipo y estado", "Perfiles de comportamiento",
                          "Casos borde (qué mirar)"):
            cell.font = Font(bold=True)
    _autowidth(ws)


def _table(ws: Worksheet, header: list[str], rows: object) -> None:
    ws.append(header)
    for row in rows:  # type: ignore[attr-defined]
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in ws.iter_cols(min_row=2):
        for cell in column:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(cell.value, time):
                cell.number_format = "hh:mm"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autowidth(ws)


def _autowidth(ws: Worksheet, sample: int = 400) -> None:
    # iter_cols crea las celdas que recorre: no pasar de la última fila con datos.
    for column in ws.iter_cols(max_row=min(sample, ws.max_row)):
        width = max((_width(c.value) for c in column), default=8)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(width + 3, 60)


def _width(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, datetime):
        return 16
    if isinstance(value, date):
        return 10
    if isinstance(value, time):
        return 5
    return len(str(value))
